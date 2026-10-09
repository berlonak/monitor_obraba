# copyright by berlonak
# telegram: @Kilax123
"""On-demand Excel export. Built in RAM; only the supervisor can request it."""
from datetime import datetime
from io import BytesIO

from .statistics import detailed_rows, overdue_rows, snapshot, work_seconds
from .timeutils import humanize_seconds, local_timezone


def make_excel(cfg, state, names, period='today', custom=None, now=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    timepoint = now or datetime.now(local_timezone(cfg['utc_offset_hours']))
    report = snapshot(state.conn, cfg, names, period, timepoint, custom)
    first, end = report['start'], report['end']
    display_names = {operator['name']: operator['operator_label'] for operator in cfg['sessions']}
    tz = local_timezone(cfg['utc_offset_hours'])

    def local(ts):
        return datetime.fromtimestamp(ts, tz).strftime('%d.%m.%Y %H:%M:%S') if ts is not None else ''

    wb = Workbook()
    summary = wb.active
    summary.title = 'Сводка'
    summary.append(['Показатель', 'Значение'])
    summary.append(['Период', report['label']])
    summary.append(['Часовой пояс', f"UTC{cfg['utc_offset_hours']:+d}"])
    summary.append(['Новые обращения', report['answers']['new']])
    summary.append(['Новые лиды', report['answers']['clients']])
    summary.append(['Ответов за период', report['answers']['handled']])
    summary.append(['Ответов в пределах SLA', report['answers']['on_time']])
    summary.append(['Нарушений SLA ответа', report['answers']['breaches']])
    summary.append(['Среднее время ответа, сек.', report['answers']['avg']])
    summary.append(['Медиана ответа, сек.', report['answers']['median']])
    summary.append(['P90 ответа, сек.', report['answers']['p90']])
    summary.append(['Среднее ответа без времени сна, сек.', report['answers']['avg_work']])
    summary.append(['Новых периодов непрочитанного', report['reads']['new']])
    summary.append(['Прочитано за период', report['reads']['handled']])
    summary.append(['Нарушений SLA прочтения', report['reads']['breaches']])
    summary.append(['Среднее время до прочтения, сек.', report['reads']['avg']])
    summary.append(['Сейчас без ответа', report['live']['answers']])
    summary.append(['Из них просрочены', report['live']['answers_overdue']])
    summary.append(['Сейчас непрочитано', report['live']['reads']])
    summary.append(['Из них просрочены', report['live']['reads_overdue']])
    summary.append(['Примечание', 'Прочтение фиксируется при опросе, точное время неизвестно.'])

    ops = wb.create_sheet('Операторы')
    ops.append(['Оператор', 'Новые обращения', 'Ответов', 'SLA ответа, %', 'Нарушений ответа',
                'Среднее ответа, сек.', 'Новых непрочитанных периодов', 'Прочитано',
                'Нарушений прочтения', 'Сейчас без ответа', 'Просрочено ответов',
                'Сейчас непрочитано', 'Просрочено непрочитанных'])
    for name in names:
        row = snapshot(state.conn, cfg, [name], period, timepoint, custom)
        a, r, live = row['answers'], row['reads'], row['live']
        ops.append([display_names.get(name, name), a['new'], a['handled'],
                    round(100 * a['on_time'] / a['handled'], 1) if a['handled'] else None,
                    a['breaches'], a['avg'], r['new'], r['handled'], r['breaches'],
                    live['answers'], live['answers_overdue'], live['reads'], live['reads_overdue']])

    overdue = wb.create_sheet('Просроченные сейчас')
    overdue.append(['Оператор', 'Telegram ID лида', f"Без ответа с (UTC{cfg['utc_offset_hours']:+d})",
                    f"Непрочитано с (UTC{cfg['utc_offset_hours']:+d})",
                    'Просрочено по ответу', 'Просрочено по прочтению'])
    current = int(timepoint.timestamp())
    for row in overdue_rows(state.conn, cfg, names, timepoint, limit=100000):
        answer, read = row['answer_ts'], row['read_ts']
        overdue.append([display_names.get(row['session'], row['session']), row['chat_id'],
                        local(answer), local(read),
                        humanize_seconds(current - answer) if answer is not None and
                        current - answer >= cfg['sla_seconds'] else '',
                        humanize_seconds(current - read) if read is not None and
                        current - read >= cfg['sla_seconds'] else ''])

    for title, kind in [('Ответы', 'answers'), ('Прочтение', 'reads')]:
        sheet = wb.create_sheet(title)
        sheet.append(['Оператор', 'Telegram ID лида', 'ID первого сообщения',
                      f"Получено (UTC{cfg['utc_offset_hours']:+d})",
                      f"Закрыто (UTC{cfg['utc_offset_hours']:+d})", 'Время, сек.',
                      'Без времени сна, сек.', 'SLA'])
        for session, chat_id, mid, start, finish, withdrawn in detailed_rows(
                state.conn, 'stats_' + kind, names, first, end):
            if withdrawn:
                # Overdue, then the message was deleted / the chat became unavailable:
                # a breach, but not an answer or a read — no response time.
                sheet.append([display_names.get(session, session), chat_id, mid, local(start),
                              local(finish), None, None, 'Нарушен; сообщение удалено или чат недоступен'])
                continue
            duration = finish - start if finish is not None else None
            sheet.append([display_names.get(session, session), chat_id, mid, local(start),
                          local(finish), duration,
                          work_seconds(start, finish, cfg.get('quiet_window'), cfg['utc_offset_hours'])
                          if finish is not None else None,
                          'Нарушен' if min(finish if finish is not None else current, current) - start >
                          cfg['sla_seconds'] else 'В пределах' if finish is not None else 'Ожидает'])

    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(fill_type='solid', fgColor='243447')
            cell.alignment = Alignment(wrap_text=True)
        for col in sheet.columns:
            letters = get_column_letter(col[0].column)
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=12)
            sheet.column_dimensions[letters].width = min(46, max(16, max_len + 2))
    data = BytesIO()
    wb.save(data)
    return f"sla_stats_{timepoint.strftime('%Y%m%d_%H%M')}.xlsx", data.getvalue()
