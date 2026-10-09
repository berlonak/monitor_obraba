# copyright by berlonak
# telegram: @Kilax123
"""Offline end-to-end stats and supervisor-only Telegram handlers."""
import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from openpyxl import load_workbook

from sla.config import parse_quiet_hours
from sla.monitor import process_dialog
from sla.state import State
from sla.statistics import (open_case, close_case, parse_custom_dates, period_bounds,
                            overdue_rows, snapshot, work_seconds)
from sla.stats_bot import StatsBot, overdue_view, summary_view
from sla.stats_excel import make_excel
from sla.timeutils import local_timezone
from test_monitor import BASE, FakeClient, FakeMessage, FakeNotifier, OP, fake_cfg, fake_dialog


class FakeBot(FakeNotifier):
    def __init__(self):
        super().__init__()
        self.requests = []
        self.documents = []

    async def request(self, name, payload):
        self.requests.append((name, payload))
        return {'ok': True, 'result': {'message_id': 11}}

    async def send_document(self, uid, filename, data):
        self.documents.append((uid, filename, data))
        return True


def config(**kwargs):
    cfg = fake_cfg(stats_enabled=True, alert_mode='unread')
    cfg['sessions'] = [OP, {'name': 'petr', 'operator_label': 'Пётр', 'operator_telegram_id': 444}]
    cfg.update(kwargs)
    return cfg


class StatisticsTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = TemporaryDirectory()
        self.db = Path(self.tmp.name) / 'state.sqlite3'
        self.state = State(self.db)
        self.bot = FakeNotifier()
        self.cfg = config()

    async def asyncTearDown(self):
        self.state.close()
        self.tmp.cleanup()

    def report(self, now=BASE, op=('ivan',)):
        return snapshot(self.state.conn, self.cfg, list(op), now=now)

    async def test_unread_only_mode_still_records_answer_time_and_read(self):
        initial = FakeMessage(1, 90)
        client = FakeClient([initial])
        await process_dialog(self.cfg, OP, client, fake_dialog([initial], 1, 0),
                             self.state, self.bot, 222, BASE)
        initial_snapshot = self.report()
        self.assertEqual(initial_snapshot['answers']['new'], 1)
        self.assertEqual(initial_snapshot['reads']['new'], 1)
        self.assertEqual(initial_snapshot['live']['answers'], 1)
        self.assertEqual(initial_snapshot['live']['reads'], 1)
        self.assertEqual(initial_snapshot['answers']['breaches'], 1)
        self.assertEqual(initial_snapshot['reads']['breaches'], 1)
        self.assertEqual(len(self.bot.deliveries), 2)  # only unread alerts, not unanswered

        reply = FakeMessage(2, 0, out=True, base=BASE + timedelta(minutes=20))
        client.messages.append(reply)
        await process_dialog(self.cfg, OP, client, fake_dialog(client.messages, 0, 1),
                             self.state, self.bot, 222, BASE + timedelta(minutes=20))
        result = self.report(BASE + timedelta(minutes=21))
        self.assertEqual(result['answers']['handled'], 1)
        self.assertEqual(result['answers']['avg'], 110 * 60)
        self.assertEqual(result['answers']['median'], 110 * 60)
        self.assertEqual(result['answers']['p90'], 110 * 60)
        self.assertEqual(result['reads']['handled'], 1)
        self.assertEqual(result['live']['answers'], 0)
        self.assertEqual(result['live']['reads'], 0)
        self.assertEqual(result['answers']['breaches'], 1)
        self.assertEqual(result['reads']['breaches'], 1)
        # Repeated unchanged polls never add duplicate cases / notifications.
        await process_dialog(self.cfg, OP, client, fake_dialog(client.messages, 0, 1),
                             self.state, self.bot, 222, BASE + timedelta(minutes=25))
        self.assertEqual(self.report(BASE + timedelta(minutes=25))['answers']['new'], 1)

    async def test_two_answer_cycles_between_polls_are_preserved(self):
        sent = FakeMessage(1, 130, out=True)
        client = FakeClient([sent])
        await process_dialog(self.cfg, OP, client, fake_dialog([sent], 0, 0),
                             self.state, self.bot, 222, BASE)
        second = [FakeMessage(2, 90), FakeMessage(3, 50, out=True), FakeMessage(4, 20)]
        client.messages += second
        await process_dialog(self.cfg, OP, client, fake_dialog(client.messages, 1, 3),
                             self.state, self.bot, 222, BASE)
        values = self.report()
        self.assertEqual(values['answers']['new'], 2)
        self.assertEqual(values['answers']['handled'], 1)
        self.assertEqual(values['answers']['avg'], 40 * 60)
        self.assertEqual(values['live']['answers'], 1)
        self.assertEqual(values['reads']['new'], 1)

    async def test_stats_survive_restart_without_repeat_counts(self):
        msg = FakeMessage(1, 90)
        client = FakeClient([msg])
        await process_dialog(self.cfg, OP, client, fake_dialog([msg], 1, 0),
                             self.state, self.bot, 222, BASE)
        self.state.close()
        self.state = State(self.db)
        await process_dialog(self.cfg, OP, client, fake_dialog([msg], 1, 0),
                             self.state, self.bot, 222, BASE + timedelta(minutes=3))
        vals = self.report(BASE + timedelta(minutes=3))
        self.assertEqual((vals['answers']['new'], vals['reads']['new']), (1, 1))

    async def test_old_v04_active_case_is_backfilled_without_data_loss(self):
        self.cfg['stats_enabled'] = False
        msg = FakeMessage(1, 90)
        client = FakeClient([msg])
        await process_dialog(self.cfg, OP, client, fake_dialog([msg], 1, 0),
                             self.state, self.bot, 222, BASE)
        self.assertEqual(self.report()['answers']['new'], 0)
        self.cfg['stats_enabled'] = True
        await process_dialog(self.cfg, OP, client, fake_dialog([msg], 1, 0),
                             self.state, self.bot, 222, BASE + timedelta(minutes=2))
        self.assertEqual(self.report(BASE + timedelta(minutes=2))['answers']['new'], 1)
        self.assertEqual(self.report(BASE + timedelta(minutes=2))['reads']['new'], 1)

    async def test_start_only_new_ignores_old_dialog_in_statistics(self):
        self.cfg['alert_existing_chats_on_start'] = False
        old = FakeMessage(1, 90)
        client = FakeClient([old])
        await process_dialog(self.cfg, OP, client, fake_dialog([old], 1, 0),
                             self.state, self.bot, 222, BASE)
        self.assertEqual(self.report()['answers']['new'], 0)
        self.assertEqual(self.report()['reads']['new'], 0)
        client.messages.append(FakeMessage(2, 30))
        await process_dialog(self.cfg, OP, client, fake_dialog(client.messages, 2, 0),
                             self.state, self.bot, 222, BASE)
        self.assertEqual(self.report()['answers']['new'], 1)
        self.assertEqual(self.report()['reads']['new'], 1)

    async def test_no_double_count_when_mode_both_sends_four_alerts(self):
        self.cfg['alert_mode'] = 'both'
        msg = FakeMessage(1, 90)
        client = FakeClient([msg])
        await process_dialog(self.cfg, OP, client, fake_dialog([msg], 1, 0),
                             self.state, self.bot, 222, BASE)
        self.assertEqual(len(self.bot.deliveries), 4)
        report = self.report()
        self.assertEqual(report['answers']['new'], 1)
        self.assertEqual(report['reads']['new'], 1)
        self.assertEqual(len(overdue_rows(self.state.conn, self.cfg, ['ivan'], BASE)), 1)

    async def test_stats_are_isolated_by_operator(self):
        open_case(self.state.conn, 'answers', 'ivan', 333, 1, int(BASE.timestamp()) - 4000)
        open_case(self.state.conn, 'answers', 'petr', 555, 1, int(BASE.timestamp()) - 3000)
        self.assertEqual(self.report()['live']['answers'], 1)
        self.assertEqual(self.report(op=('petr',))['live']['answers'], 1)
        self.assertEqual(self.report(op=('ivan', 'petr'))['live']['answers'], 2)
        self.assertEqual(len(overdue_rows(self.state.conn, self.cfg, ['ivan', 'petr'], BASE)), 1)

    async def test_periods_and_quiet_adjustment(self):
        tz = local_timezone(3)
        start = datetime(2026, 9, 24, 22, 30, tzinfo=tz)
        stop = datetime(2026, 9, 25, 10, 30, tzinfo=tz)
        quiet = parse_quiet_hours({'start': '23:00', 'end': '09:00'})
        self.assertEqual(work_seconds(int(start.timestamp()), int(stop.timestamp()), quiet, 3), 7200)
        self.assertEqual(work_seconds(int(start.timestamp()), int(stop.timestamp()), None, 3), 43200)
        label, lo, hi = period_bounds('yesterday', 3, now=BASE)
        self.assertEqual(label, 'вчера')
        self.assertLess(lo, hi)
        self.assertEqual(parse_custom_dates(['01.09.2026', '25.09.2026']),
                         (datetime(2026, 9, 1).date(), datetime(2026, 9, 25).date()))
        with self.assertRaises(ValueError):
            period_bounds('today', 3, custom=parse_custom_dates(['01.09.2024', '01.09.2026']))

    async def test_excel_has_summary_operator_cases_and_backlog(self):
        open_case(self.state.conn, 'answers', 'ivan', 333, 1, int(BASE.timestamp()) - 6000)
        open_case(self.state.conn, 'reads', 'ivan', 333, 1, int(BASE.timestamp()) - 6000)
        close_case(self.state.conn, 'reads', 'ivan', 333, 1, int(BASE.timestamp()) - 100)
        filename, content = make_excel(self.cfg, self.state, ['ivan'], now=BASE)
        self.assertTrue(filename.endswith('.xlsx'))
        wb = load_workbook(BytesIO(content), read_only=True)
        self.assertEqual(wb.sheetnames, ['Сводка', 'Операторы', 'Просроченные сейчас', 'Ответы', 'Прочтение'])
        self.assertEqual(wb['Просроченные сейчас']['B2'].value, 333)
        self.assertEqual(wb['Операторы']['A2'].value, 'Иван')
        wb.close()

    async def test_report_text_and_overdue_links(self):
        open_case(self.state.conn, 'answers', 'ivan', 333, 1, int(BASE.timestamp()) - 7200)
        text, buttons = summary_view(self.cfg, self.state, now=BASE)
        self.assertIn('Все операторы', text)
        self.assertIn('просрочено', text)
        self.assertIn('Иван', str(buttons))
        late, links = overdue_view(self.cfg, self.state, now=BASE)
        self.assertIn('tg://user?id=333', late)
        self.assertIn('лид 333', late)


class StatsBotAuthTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = TemporaryDirectory()
        self.state = State(Path(self.tmp.name) / 'state.db')
        self.cfg = config()
        self.bot = FakeBot()
        self.handler = StatsBot(self.cfg, self.state, self.bot)

    async def asyncTearDown(self):
        self.state.close()
        self.tmp.cleanup()

    async def test_operator_and_groups_get_no_stats(self):
        for sender_id, chat_id, chat_type in ((222, 222, 'private'), (111, -10011, 'supergroup'),
                                              (222, 111, 'private'), (111, 222, 'private')):
            await self.handler.on_message({'text': '/stats', 'from': {'id': sender_id},
                                           'chat': {'id': chat_id, 'type': chat_type}})
        self.assertEqual(self.bot.requests, [])
        self.assertEqual(self.bot.deliveries, [])
        await self.handler.on_callback({'id': 'cb', 'data': 'p:today', 'from': {'id': 222},
                                        'message': {'chat': {'id': 222, 'type': 'private'}}})
        self.assertEqual([x[0] for x in self.bot.requests], ['answerCallbackQuery'])
        self.assertEqual(self.bot.documents, [])

    async def test_supervisor_can_open_period_operator_overdue_and_excel(self):
        msg = lambda text: {'text': text, 'from': {'id': 111},
                            'chat': {'id': 111, 'type': 'private'}}
        callback = lambda data: {'id': 'cb', 'data': data, 'from': {'id': 111},
                                 'message': {'chat': {'id': 111, 'type': 'private'}, 'message_id': 2}}
        await self.handler.on_message(msg('/stats'))
        await self.handler.on_callback(callback('p:week'))
        await self.handler.on_callback(callback('o:today:0'))
        await self.handler.on_callback(callback('l:a:0'))
        await self.handler.on_callback(callback('x:today:0'))
        await self.handler.on_message(msg('/stats_excel 01.09.2026 24.09.2026'))
        await self.handler.on_message(msg('/stats_excel ivan 01.09.2026 24.09.2026'))
        self.assertEqual(len(self.bot.documents), 3)
        self.assertEqual(set(x[0] for x in self.bot.documents), {111})
        self.assertEqual(sum(action == 'editMessageText' for action, _ in self.bot.requests), 3)
        self.assertEqual(sum(action == 'sendMessage' for action, _ in self.bot.requests), 1)

    async def test_invalid_dates_cannot_crash_bot(self):
        await self.handler.on_message({'text': '/stats 01.01.2000 01.09.2026',
                                       'from': {'id': 111},
                                       'chat': {'id': 111, 'type': 'private'}})
        self.assertIn('Период', self.bot.deliveries[0][1])
