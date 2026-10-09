# copyright by berlonak
# telegram: @Kilax123
"""v0.5.5: developer-only technical alerts, reconnects, delivery and lead-logic fixes."""
import asyncio
from collections import Counter
from datetime import timedelta
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import yaml

from sla.config import load_config
from sla.devalerts import TechAlerts, supervise
from sla.instance_lock import AlreadyRunning, InstanceLock
from sla.monitor import (_check_archive_setting, _reconcile, dialog_skip_reason,
                         monitor_session, process_dialog)
from sla.notifier import Delivery, Notifier
from sla.resolutions import apply_pending_edits
from sla.state import State, ChatState, UnreadState
from test_monitor import BASE, FakeClient, FakeMessage, OP, fake_cfg, fake_dialog
from test_resolutions import TrackedBot

DEV, MAIN, OPERATOR, LEAD_ID = 999, 111, 222, 333
DAY = 24 * 60


class LookupClient(FakeClient):
    """FakeClient that also answers get_messages(ids=...) like Telethon (None if deleted)."""
    def __init__(self, messages):
        super().__init__(messages)
        self.lookups = []

    async def get_messages(self, entity, ids=None):
        self.lookups.append(ids)
        return next((m for m in self.messages if m.id == ids), None)


class DeliveryBot:
    """Notifier double with the real deliver() contract."""
    def __init__(self, permanent=(), transient=()):
        self.permanent = set(permanent)
        self.transient = set(transient)
        self.sent = []
        self.edits = []
        self.next_id = 500

    async def deliver(self, recipient, body):
        if recipient in self.permanent:
            return Delivery("permanent", None, "Forbidden: bot was blocked by the user")
        if recipient in self.transient:
            return Delivery("transient", None, "timeout")
        self.next_id += 1
        self.sent.append((recipient, body))
        return Delivery("ok", self.next_id, None)

    async def send(self, recipient, body):
        return (await self.deliver(recipient, body)).status == "ok"

    async def edit(self, recipient, message_id, body):
        self.edits.append((recipient, message_id, body))
        return "edited", None

    def to(self, recipient):
        return [body for who, body in self.sent if who == recipient]


def dialog_for(uid, messages, unread_count=0, read_max_id=None):
    dialog = fake_dialog(messages, unread_count=unread_count, read_max_id=read_max_id)
    dialog.entity = SimpleNamespace(id=uid, first_name="Лид", last_name=str(uid), username=None,
                                    bot=False, deleted=False, contact=False)
    return dialog


def dev_cfg(**overrides):
    return fake_cfg(developer_telegram_id=DEV, **overrides)


class StateCase(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dbfile = Path(self.tmp.name) / "state.sqlite3"
        self.state = State(self.dbfile)

    async def asyncTearDown(self):
        self.state.close()
        self.tmp.cleanup()


# ---------------------------------------------------------------- developer channel

class TechAlertsTests(unittest.IsolatedAsyncioTestCase):
    async def test_dedup_repeat_resolve_only_to_developer(self):
        now = [1000.0]
        bot = DeliveryBot()
        tech = TechAlerts(dev_cfg(tech_alert_repeat_seconds=3600), bot, clock=lambda: now[0])
        self.assertTrue(await tech.problem("link:ivan", "нет связи"))
        self.assertFalse(await tech.problem("link:ivan", "нет связи"))
        now[0] += 3599
        self.assertFalse(await tech.problem("link:ivan", "нет связи"))
        now[0] += 1
        self.assertTrue(await tech.problem("link:ivan", "нет связи"))
        self.assertIn("Напоминание", bot.to(DEV)[-1])
        self.assertTrue(await tech.resolved("link:ivan", "связь есть"))
        self.assertFalse(await tech.resolved("link:ivan", "связь есть"))
        self.assertEqual(len(bot.to(DEV)), 3)
        self.assertEqual([who for who, _ in bot.sent], [DEV, DEV, DEV])

    async def test_no_repeat_flag_and_no_developer_means_log_only(self):
        now = [0.0]
        bot = DeliveryBot()
        tech = TechAlerts(dev_cfg(tech_alert_repeat_seconds=60), bot, clock=lambda: now[0])
        await tech.problem("archive:ivan", "автоархив", repeat=False)
        now[0] += 10 ** 6
        await tech.problem("archive:ivan", "автоархив", repeat=False)
        self.assertEqual(len(bot.sent), 1)
        silent = TechAlerts(fake_cfg(), bot)
        with self.assertLogs("sla.devalerts", "WARNING"):
            self.assertFalse(await silent.problem("link:ivan", "нет связи"))
        self.assertFalse(await silent.resolved("link:ivan", "ok"))
        self.assertEqual(len(bot.sent), 1)

    async def test_transient_failures_are_retried_and_resolved_needs_delivered_problem(self):
        waits = []

        async def fake_sleep(seconds):
            waits.append(seconds)
        bot = DeliveryBot(transient=[DEV])
        tech = TechAlerts(dev_cfg(), bot)
        with patch("sla.devalerts.asyncio.sleep", fake_sleep), self.assertLogs("sla.devalerts", "WARNING"):
            self.assertFalse(await tech.problem("k", "сбой"))
            self.assertEqual(waits, [2, 5])  # three attempts before giving up for now
            self.assertFalse(await tech.resolved("k", "ok"))  # developer never saw it: silent
            self.assertFalse(await tech.problem("k2", "сбой"))
        bot.transient.clear()
        self.assertTrue(await tech.problem("k2", "сбой"))  # next report of the same problem retries
        self.assertTrue(await tech.resolved("k2", "ok"))
        self.assertEqual(len(bot.to(DEV)), 2)

    async def test_secrets_are_redacted(self):
        bot = DeliveryBot()
        cfg = dev_cfg(notifier_bot_token="123:SECRETTOKENSECRETTOKENSECRET", api_hash="b" * 32)
        await TechAlerts(cfg, bot).info("url https://api.telegram.org/bot123:SECRETTOKENSECRETTOKENSECRET/x "
                                        + "b" * 32)
        body = bot.to(DEV)[0]
        self.assertNotIn("SECRETTOKEN", body)
        self.assertNotIn("b" * 32, body)

    async def test_supervise_reports_crash_and_restarts(self):
        bot = DeliveryBot()
        tech = TechAlerts(dev_cfg(), bot)
        runs = []

        async def flaky():
            runs.append(1)
            if len(runs) == 1:
                raise RuntimeError("boom")

        async def no_wait(_):
            return None
        with patch("sla.devalerts.asyncio.sleep", no_wait), self.assertLogs("sla.devalerts", "ERROR"):
            await supervise("бот статистики", flaky, tech, restart_delay=1)
        self.assertEqual(len(runs), 2)
        self.assertEqual(len(bot.to(DEV)), 1)
        self.assertIn("бот статистики", bot.to(DEV)[0])
        self.assertIn("RuntimeError: boom", bot.to(DEV)[0])


class FakeErrors:
    class FloodWaitError(Exception):
        def __init__(self, seconds):
            super().__init__(seconds)
            self.seconds = seconds

    class UnauthorizedError(Exception):
        pass

    class AuthKeyError(Exception):
        pass

    class AuthKeyDuplicatedError(AuthKeyError):
        pass


def make_telethon(client_cls):
    return SimpleNamespace(TelegramClient=client_cls, errors=FakeErrors)


class SessionSupervisionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        path = Path(self.tmp.name) / "operator.session"
        path.touch()
        self.sess = dict(OP, session_file=str(path.with_suffix("")))
        self.state = State(Path(self.tmp.name) / "state.sqlite3")
        self.bot = DeliveryBot()
        self.sleeps = []

    async def asyncTearDown(self):
        self.state.close()
        self.tmp.cleanup()

    def cfg(self, **extra):
        return dict(dev_cfg(history_request_delay_seconds=0, poll_interval_seconds=123, **extra),
                    api_id=1, api_hash="a" * 32, include_archived=True, flood_wait_extra_seconds=3)

    async def run_session(self, client_cls, cfg, stop_after_poll_sleeps=1):
        polls = []

        async def fake_sleep(seconds):
            self.sleeps.append(seconds)
            if seconds == cfg["poll_interval_seconds"]:
                polls.append(seconds)
                if len(polls) >= stop_after_poll_sleeps:
                    raise asyncio.CancelledError
        tech = TechAlerts(cfg, self.bot)
        with patch.dict(sys.modules, {"telethon": make_telethon(client_cls)}), \
                patch("sla.monitor.asyncio.sleep", fake_sleep):
            try:
                await monitor_session(cfg, self.sess, self.state, self.bot, tech)
                return "returned"
            except asyncio.CancelledError:
                return "cancelled"

    def assert_nothing_to_staff(self):
        self.assertEqual(self.bot.to(MAIN), [])
        self.assertEqual(self.bot.to(OPERATOR), [])

    async def test_network_down_at_start_keeps_retrying_and_tells_only_developer(self):
        class Client:
            instances = connects = 0
            def __init__(self, *args):
                type(self).instances += 1
            async def connect(self):
                type(self).connects += 1
                if type(self).connects <= 2:
                    raise ConnectionError("network is unreachable")
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            result = await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0))
        self.assertEqual(result, "cancelled")  # never exits on its own
        self.assertEqual(Client.connects, 3)
        self.assertEqual(Client.instances, 3)  # a fresh client after each failure
        dev = self.bot.to(DEV)
        self.assertEqual(len(dev), 2)
        self.assertIn("нет связи с Telegram", dev[0])
        self.assertIn("связь с Telegram восстановлена", dev[1])
        self.assertEqual(self.sleeps[:2], [5, 10])  # backoff between attempts
        self.assert_nothing_to_staff()

    async def test_connection_lost_mid_run_reconnects_short_blip_is_silent(self):
        class Client:
            instances = passes = 0
            def __init__(self, *args):
                type(self).instances += 1
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                type(self).passes += 1
                if type(self).passes == 1:
                    raise ConnectionError("Cannot send requests while disconnected")
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            result = await self.run_session(Client, self.cfg(tech_alert_delay_seconds=300))
        self.assertEqual(result, "cancelled")
        self.assertEqual(Client.instances, 2)
        self.assertEqual(Client.passes, 2)
        self.assertEqual(self.bot.sent, [])  # shorter than tech_alert_delay_seconds

    async def test_revoked_session_stops_operator_and_tells_developer(self):
        class Client:
            def __init__(self, *args): pass
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                raise FakeErrors.UnauthorizedError("AUTH_KEY_UNREGISTERED")
                yield

        with self.assertLogs("sla.monitor", "ERROR"):
            result = await self.run_session(Client, self.cfg())
        self.assertEqual(result, "returned")
        self.assertEqual(len(self.bot.to(DEV)), 1)
        self.assertIn("отклонил сессию", self.bot.to(DEV)[0])
        self.assert_nothing_to_staff()

    async def test_unauthorized_and_missing_session_are_reported(self):
        class Client:
            def __init__(self, *args): pass
            async def connect(self): pass
            async def is_user_authorized(self): return False
            async def get_me(self): return None
            async def disconnect(self): pass

        with self.assertLogs("sla.monitor", "ERROR"):
            self.assertEqual(await self.run_session(Client, self.cfg()), "returned")
        self.assertIn("не авторизована", self.bot.to(DEV)[0])
        missing = dict(self.sess, session_file=str(Path(self.tmp.name) / "nobody"))
        tech = TechAlerts(self.cfg(), self.bot)
        with patch.dict(sys.modules, {"telethon": make_telethon(Client)}), \
                self.assertLogs("sla.monitor", "ERROR"):
            await monitor_session(self.cfg(), missing, self.state, self.bot, tech)
        self.assertIn("нет файла сессии", self.bot.to(DEV)[1])
        self.assert_nothing_to_staff()

    async def test_floodwait_masked_as_unauthorized_is_retried(self):
        class Client:
            calls = 0
            def __init__(self, *args): pass
            async def connect(self): pass
            async def is_user_authorized(self): return False  # Telethon: any RPC error -> False
            async def get_me(self):
                type(self).calls += 1
                if type(self).calls == 1:
                    raise FakeErrors.FloodWaitError(4)
                return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            self.assertEqual(await self.run_session(Client, self.cfg()), "cancelled")
        self.assertEqual(self.sleeps[0], 7)  # FloodWait 4 + 3 extra, not "unauthorized"
        self.assertEqual(self.bot.sent, [])

    async def test_vanished_chat_is_forgotten_after_three_complete_passes(self):
        self.state.save("ivan", LEAD_ID, ChatState(last_seen_id=5, first_unanswered_id=5,
                                                    first_unanswered_ts=100))
        self.state.conn.execute("INSERT INTO stats_answers (session, chat_id, message_id, opened_ts, closed_ts) VALUES ('ivan', ?, 5, 100, NULL)", (LEAD_ID,))
        self.state.conn.commit()

        class Client:
            def __init__(self, *args): pass
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        await self.run_session(Client, self.cfg(), stop_after_poll_sleeps=2)
        self.assertIsNotNone(self.state.get("ivan", LEAD_ID))  # two passes: still kept
        with self.assertLogs("sla.monitor", "INFO"):
            await self.run_session(Client, self.cfg(), stop_after_poll_sleeps=3)
        self.assertIsNone(self.state.get("ivan", LEAD_ID))
        self.assertEqual(self.state.conn.execute(
            "SELECT count(*) FROM stats_answers WHERE closed_ts IS NULL").fetchone()[0], 0)

    async def test_archive_autoarchive_warning_goes_to_developer(self):
        class Client:
            async def __call__(self, request):
                return SimpleNamespace(archive_and_mute_new_noncontact_peers=True)
        tech = TechAlerts(dev_cfg(), self.bot)
        with self.assertLogs("sla.monitor", "WARNING"):
            await _check_archive_setting(Client(), dict(fake_cfg(), include_archived=False), "ivan", "Оператор", tech)
        self.assertIn("include_archived", self.bot.to(DEV)[0])
        await _check_archive_setting(Client(), dict(fake_cfg(), include_archived=True), "ivan", "Оператор", tech)
        self.assertEqual(len(self.bot.sent), 1)


# ---------------------------------------------------------------- delivery

class DeliveryTests(StateCase):
    async def test_blocked_operator_does_not_stop_supervisor_reminders(self):
        cfg = dev_cfg(alert_mode="unanswered", remind_every_seconds=600)
        bot = DeliveryBot(permanent=[OPERATOR])
        tech = TechAlerts(cfg, bot)
        msgs = [FakeMessage(1, 90)]
        for minutes in (0, 10, 20, 30):
            await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR,
                                 BASE + timedelta(minutes=minutes), tech=tech)
        self.assertEqual(len(bot.to(MAIN)), 4)  # first alert + 3 reminders
        self.assertEqual(self.state.get("ivan", LEAD_ID).reminders_sent, 3)
        self.assertEqual(len(bot.to(DEV)), 1)  # told once, not every pass
        self.assertIn("оператору «Иван»", bot.to(DEV)[0])
        self.assertIn("bot was blocked", bot.to(DEV)[0])
        bot.permanent.clear()
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR,
                             BASE + timedelta(minutes=40), tech=tech)
        self.assertEqual(len(bot.to(OPERATOR)), 1)
        self.assertIn("снова доставляет", bot.to(DEV)[-1])

    async def test_transient_failure_is_still_retried_next_pass(self):
        cfg = dev_cfg(alert_mode="unanswered")
        bot = DeliveryBot(transient=[MAIN])
        tech = TechAlerts(cfg, bot)
        msgs = [FakeMessage(1, 90)]
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR, BASE, tech=tech)
        self.assertTrue(self.state.get("ivan", LEAD_ID).pending_main)
        bot.transient.clear()
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR,
                             BASE + timedelta(minutes=1), tech=tech)
        self.assertEqual(len(bot.to(MAIN)), 1)
        self.assertEqual(bot.to(DEV), [])


class NotifierTests(unittest.IsolatedAsyncioTestCase):
    def bot(self):
        class Bot(Notifier):
            async def request(self, method, data=None):
                self.calls.append(method)
                return self.reply
        bot = Bot({"notifier_bot_token": "test"})
        bot.calls = []
        return bot

    async def test_error_classification_and_429_pause(self):
        bot = self.bot()
        for code, text, expected in ((403, "Forbidden: bot was blocked by the user", "permanent"),
                                     (400, "Bad Request: chat not found", "permanent"),
                                     (500, "Internal Server Error", "transient")):
            bot.reply = {"ok": False, "error_code": code, "description": text}
            with self.assertLogs("sla.notifier", "ERROR"):
                result = await bot.deliver(1, "x")
            self.assertEqual((result.status, result.description), (expected, text))
        bot.reply = {"ok": False, "error_code": 429, "description": "Too Many Requests",
                     "parameters": {"retry_after": 7}}
        with self.assertLogs("sla.notifier", "WARNING"):
            self.assertEqual((await bot.deliver(1, "x")).status, "transient")
        calls = len(bot.calls)
        self.assertEqual((await bot.deliver(1, "x")).status, "transient")
        self.assertEqual(len(bot.calls), calls)  # paused: Bot API not hammered
        bot.reply = {"ok": True, "result": {"message_id": 5}}
        bot._send_paused_until = 0
        self.assertEqual(await bot.send_tracked(1, "x"), (True, 5))

    async def test_edit_to_blocked_chat_is_terminal(self):
        bot = self.bot()
        bot.reply = {"ok": False, "error_code": 403, "description": "Forbidden: bot was blocked by the user"}
        with self.assertLogs("sla.notifier", "INFO"):
            self.assertEqual(await bot.edit(1, 2, "x"), ("gone", None))

    async def test_start_waits_for_network_but_rejects_bad_token(self):
        attempts = []

        class Bot(Notifier):
            async def request(self, method, data=None):
                attempts.append(method)
                if len(attempts) < 3:
                    raise OSError("network is unreachable")
                return self.reply

        async def no_wait(_):
            return None
        bot = Bot({"notifier_bot_token": "test"})
        bot.reply = {"ok": True, "result": {"username": "sla_bot"}}
        with patch("sla.notifier.asyncio.sleep", no_wait), self.assertLogs("sla.notifier", "WARNING"):
            await bot.start()
        self.assertEqual(len(attempts), 3)
        await bot.stop()
        bad = Bot({"notifier_bot_token": "test"})
        bad.reply = {"ok": False, "error_code": 401, "description": "Unauthorized"}
        attempts.clear()
        attempts.extend([1, 2])
        with self.assertRaisesRegex(ValueError, "notifier_bot_token"):
            await bad.start()


# ---------------------------------------------------------------- lead logic

class FilterTests(unittest.TestCase):
    def test_service_account_and_staff_are_not_leads(self):
        cfg = dev_cfg(sessions=[{"operator_telegram_id": 444}])
        msgs = [FakeMessage(1, 5)]
        self.assertEqual(dialog_skip_reason(dialog_for(777000, msgs), cfg, OPERATOR), "служебные_telegram")
        support = dialog_for(12345, msgs)
        support.entity.support = True
        self.assertEqual(dialog_skip_reason(support, cfg, OPERATOR), "служебные_telegram")
        for staff in (MAIN, DEV, 444):
            self.assertEqual(dialog_skip_reason(dialog_for(staff, msgs), cfg, OPERATOR), "сотрудники")
        self.assertIsNone(dialog_skip_reason(dialog_for(LEAD_ID, msgs), cfg, OPERATOR))

    def test_config_staff_ids_and_new_settings(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "config.yaml"
            base = {"api_id": 1, "api_hash": "a" * 32, "notifier_bot_token": "1234:" + "A" * 30,
                    "main_telegram_id": MAIN,
                    "sessions": [{"name": "op", "operator_telegram_id": OPERATOR, "session_file": "s/op"}]}
            path.write_text(yaml.safe_dump(base), encoding="utf-8")
            cfg = load_config(path)
            self.assertIsNone(cfg["developer_telegram_id"])
            self.assertEqual((cfg["history_max_age_hours"], cfg["tech_alert_delay_seconds"],
                              cfg["tech_alert_repeat_seconds"]), (72, 300, 21600))
            self.assertEqual(cfg["_staff_ids"], frozenset({MAIN, OPERATOR}))
            path.write_text(yaml.safe_dump(dict(base, developer_telegram_id=DEV)), encoding="utf-8")
            self.assertEqual(load_config(path)["_staff_ids"], frozenset({MAIN, OPERATOR, DEV}))
            for bad in ({"developer_telegram_id": "999"}, {"developer_telegram_id": 0},
                        {"history_max_age_hours": -1}, {"tech_alert_delay_seconds": True}):
                path.write_text(yaml.safe_dump(dict(base, **bad)), encoding="utf-8")
                with self.subTest(bad=bad), self.assertRaises(ValueError):
                    load_config(path)


class HistoryWindowTests(StateCase):
    async def test_dead_unanswered_chat_opens_no_case_and_costs_no_history(self):
        cfg = fake_cfg(alert_mode="unanswered", history_max_age_hours=72, stats_enabled=True)
        msgs = [FakeMessage(1, 5 * DAY)]
        client = FakeClient(msgs)
        bot = DeliveryBot()
        await process_dialog(cfg, OP, client, fake_dialog(msgs), self.state, bot, OPERATOR, BASE)
        self.assertIsNone(self.state.get("ivan", LEAD_ID).first_unanswered_id)
        self.assertEqual(client.calls, [])
        self.assertEqual(bot.sent, [])

    async def test_old_and_recent_messages_case_counts_from_window(self):
        cfg = fake_cfg(alert_mode="unanswered", history_max_age_hours=72)
        msgs = [FakeMessage(1, 5 * DAY), FakeMessage(2, 120)]
        bot = DeliveryBot()
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR, BASE)
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 2)
        self.assertIn("2 ч", bot.to(MAIN)[0])

    async def test_unread_discovery_uses_window(self):
        cfg = fake_cfg(alert_mode="unread", history_max_age_hours=72)
        msgs = [FakeMessage(1, 5 * DAY), FakeMessage(2, 90)]
        bot = DeliveryBot()
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs, 2, 0), self.state, bot, OPERATOR, BASE)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 2)
        self.assertIn("1 ч 30 мин", bot.to(MAIN)[0])

    async def test_revived_dead_unread_chat_counts_from_new_message(self):
        cfg = fake_cfg(alert_mode="unread", history_max_age_hours=72)
        msgs = [FakeMessage(i, 5 * DAY - i) for i in range(1, 4)]
        client = FakeClient(msgs)
        bot = DeliveryBot()
        await process_dialog(cfg, OP, client, fake_dialog(msgs, 3, 0), self.state, bot, OPERATOR, BASE)
        self.assertIsNone(self.state.get_unread("ivan", LEAD_ID).first_unread_id)
        self.assertEqual(client.calls, [])
        client.messages.append(FakeMessage(4, 1))
        await process_dialog(cfg, OP, client, fake_dialog(client.messages, 4, 0), self.state, bot,
                             OPERATOR, BASE)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 4)
        self.assertEqual(bot.sent, [])  # 1 minute old: no false "5 days unread" alert

    async def test_partial_read_continuation_ignores_window(self):
        cfg = fake_cfg(alert_mode="unread", history_max_age_hours=72)
        msgs = [FakeMessage(1, 80 * 60), FakeMessage(2, 79 * 60)]
        self.state.save_unread("ivan", LEAD_ID, UnreadState(
            last_seen_id=2, first_unread_id=1, first_unread_ts=int(msgs[0].date.timestamp()),
            last_read_max_id=0, last_unread_count=2, read_marker_known=True))
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs, 1, 1), self.state,
                             DeliveryBot(), OPERATOR, BASE)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 2)

    async def test_one_time_cleanup_of_ancient_cases(self):
        old = int(BASE.timestamp()) - 400 * 86400
        self.state.save("ivan", LEAD_ID, ChatState(last_seen_id=9, first_unanswered_id=9, first_unanswered_ts=old))
        self.state.save("ivan", 444, ChatState(last_seen_id=3, first_unanswered_id=3,
                                                first_unanswered_ts=int(BASE.timestamp()) - 3600))
        self.state.conn.execute("INSERT INTO stats_answers (session, chat_id, message_id, opened_ts, closed_ts) VALUES ('ivan', ?, 9, ?, NULL)", (LEAD_ID, old))
        self.state.conn.commit()
        cutoff = int(BASE.timestamp()) - 72 * 3600
        with self.assertLogs("sla.state", "INFO"):
            self.assertEqual(self.state.expire_ancient_cases(cutoff), 1)
        self.assertIsNone(self.state.get("ivan", LEAD_ID))
        self.assertIsNotNone(self.state.get("ivan", 444))
        self.assertEqual(self.state.conn.execute("SELECT count(*) FROM stats_answers").fetchone()[0], 0)
        self.state.close()
        self.state = State(self.dbfile)
        self.assertIsNone(self.state.expire_ancient_cases(cutoff))  # runs only once


class AutoReplyTests(StateCase):
    async def test_business_auto_replies_do_not_close_unanswered(self):
        cfg = fake_cfg(alert_mode="unanswered", stats_enabled=True)
        client = FakeClient([FakeMessage(1, 90)])
        bot = DeliveryBot()
        await process_dialog(cfg, OP, client, fake_dialog(client.messages), self.state, bot, OPERATOR, BASE)
        away = FakeMessage(2, 89, out=True, text="Ответим в рабочее время")
        away.offline = True
        business_bot = FakeMessage(3, 88, out=True, text="Я бот")
        business_bot.via_business_bot_id = 42
        client.messages += [away, business_bot]
        await process_dialog(cfg, OP, client, fake_dialog(client.messages), self.state, bot, OPERATOR,
                             BASE + timedelta(minutes=1))
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 1)
        client.messages.append(FakeMessage(4, 0, out=True, text="Здравствуйте!"))
        await process_dialog(cfg, OP, client, fake_dialog(client.messages), self.state, bot, OPERATOR,
                             BASE + timedelta(minutes=2))
        self.assertIsNone(self.state.get("ivan", LEAD_ID).first_unanswered_id)
        closed = self.state.conn.execute("SELECT closed_ts FROM stats_answers").fetchone()[0]
        self.assertEqual(closed, int(client.messages[-1].date.timestamp()))

    async def test_first_scan_skips_auto_reply(self):
        cfg = fake_cfg(alert_mode="unanswered")
        away = FakeMessage(2, 80, out=True)
        away.offline = True
        msgs = [FakeMessage(1, 90), away]
        await process_dialog(cfg, OP, FakeClient(msgs), fake_dialog(msgs), self.state, DeliveryBot(),
                             OPERATOR, BASE)
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 1)


class DeletedMessageTests(StateCase):
    async def test_deleted_unread_message_closes_case_and_marks_alerts(self):
        cfg = fake_cfg(alert_mode="unread", stats_enabled=True, edit_resolved_alerts=True)
        bot = TrackedBot()
        msgs = [FakeMessage(5, 300, out=True), FakeMessage(10, 90)]
        await process_dialog(cfg, OP, LookupClient(msgs), fake_dialog(msgs, 1, 5), self.state, bot, OPERATOR, BASE)
        self.assertEqual(len(bot.sent), 2)
        after = [FakeMessage(5, 300, out=True)]
        with self.assertLogs("sla.monitor", "INFO"):
            await process_dialog(cfg, OP, LookupClient(after), fake_dialog(after, 0, 5), self.state, bot,
                                 OPERATOR, BASE + timedelta(minutes=5))
        self.assertIsNone(self.state.get_unread("ivan", LEAD_ID).first_unread_id)
        self.assertEqual(self.state.conn.execute(
            "SELECT count(*) FROM stats_reads WHERE closed_ts IS NULL").fetchone()[0], 0)
        now_ts = int((BASE + timedelta(minutes=5)).timestamp())
        await apply_pending_edits(self.state, bot, cfg, now_ts=now_ts, pace_seconds=0)
        self.assertEqual(len(bot.edits), 2)
        self.assertTrue(all(body.startswith("⚪️ <b>Сообщение лида удалено из чата</b>") for _, _, body in bot.edits))
        self.assertTrue(all("лид или оператор" in body for _, _, body in bot.edits))

    async def test_deleted_unread_then_new_message_has_fresh_clock(self):
        cfg = fake_cfg(alert_mode="unread")
        bot = DeliveryBot()
        msgs = [FakeMessage(5, 300, out=True), FakeMessage(10, 30)]
        await process_dialog(cfg, OP, LookupClient(msgs), fake_dialog(msgs, 1, 5), self.state, bot, OPERATOR, BASE)
        later = BASE + timedelta(hours=3)
        new = [FakeMessage(5, 480, out=True, base=later), FakeMessage(11, 1, base=later)]
        with self.assertLogs("sla.monitor", "INFO"):
            await process_dialog(cfg, OP, LookupClient(new), fake_dialog(new, 1, 5), self.state, bot,
                                 OPERATOR, later)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 11)
        self.assertEqual(bot.sent, [])  # the old code alerted "3 h 30 min unread" here

    async def test_existing_message_with_changes_costs_one_lookup_only(self):
        cfg = fake_cfg(alert_mode="unread")
        msgs = [FakeMessage(5, 300, out=True), FakeMessage(10, 30)]
        client = LookupClient(msgs)
        await process_dialog(cfg, OP, client, fake_dialog(msgs, 1, 5), self.state, DeliveryBot(), OPERATOR, BASE)
        client.messages.append(FakeMessage(11, 1))
        await process_dialog(cfg, OP, client, fake_dialog(client.messages, 2, 5), self.state,
                             DeliveryBot(), OPERATOR, BASE)
        await process_dialog(cfg, OP, client, fake_dialog(client.messages, 2, 5), self.state,
                             DeliveryBot(), OPERATOR, BASE)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 10)
        self.assertEqual(client.lookups, [10])  # unchanged chat: no extra requests

    async def test_both_clocks_share_one_lookup(self):
        cfg = fake_cfg(alert_mode="unread", stats_enabled=True)
        msgs = [FakeMessage(5, 300, out=True), FakeMessage(10, 30)]
        client = LookupClient(msgs)
        await process_dialog(cfg, OP, client, fake_dialog(msgs, 1, 5), self.state, DeliveryBot(), OPERATOR, BASE)
        client.messages.append(FakeMessage(11, 1))
        await process_dialog(cfg, OP, client, fake_dialog(client.messages, 2, 5), self.state,
                             DeliveryBot(), OPERATOR, BASE)
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 10)
        self.assertEqual(self.state.get_unread("ivan", LEAD_ID).first_unread_id, 10)
        self.assertEqual(client.lookups, [10])

    async def test_deleted_unanswered_message_closes_case(self):
        cfg = fake_cfg(alert_mode="unanswered", stats_enabled=True, edit_resolved_alerts=True)
        bot = TrackedBot()
        msgs = [FakeMessage(5, DAY, out=True), FakeMessage(10, 90)]
        await process_dialog(cfg, OP, LookupClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR, BASE)
        self.assertEqual(len(bot.sent), 2)
        after = [FakeMessage(5, DAY, out=True)]
        with self.assertLogs("sla.monitor", "INFO"):
            await process_dialog(cfg, OP, LookupClient(after), fake_dialog(after), self.state, bot,
                                 OPERATOR, BASE + timedelta(minutes=5))
        self.assertIsNone(self.state.get("ivan", LEAD_ID).first_unanswered_id)
        # Already overdue when deleted: the breach stays in statistics (closed, not erased).
        row = self.state.conn.execute("SELECT opened_ts, closed_ts FROM stats_answers").fetchall()
        self.assertEqual(len(row), 1)
        self.assertEqual(row[0][1], int((BASE + timedelta(minutes=5)).timestamp()))
        await process_dialog(cfg, OP, LookupClient(after), fake_dialog(after), self.state, bot,
                             OPERATOR, BASE + timedelta(minutes=30))
        self.assertEqual(len(bot.sent), 2)  # no reminders for a withdrawn message
        self.assertEqual(len(self.state.pending_edits(int(BASE.timestamp()) + 3600)), 2)

    async def test_deleted_unanswered_then_new_message_reanchors_clock(self):
        cfg = fake_cfg(alert_mode="unanswered")
        bot = DeliveryBot()
        msgs = [FakeMessage(5, DAY, out=True), FakeMessage(10, 30)]
        await process_dialog(cfg, OP, LookupClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR, BASE)
        later = BASE + timedelta(hours=3)
        new = [FakeMessage(5, DAY, out=True, base=later), FakeMessage(11, 1, base=later)]
        with self.assertLogs("sla.monitor", "INFO"):
            await process_dialog(cfg, OP, LookupClient(new), fake_dialog(new), self.state, bot, OPERATOR, later)
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 11)
        self.assertEqual(bot.sent, [])


class ReconcileTests(StateCase):
    async def test_missed_counter_resets_when_chat_is_seen_again(self):
        self.state.save("ivan", LEAD_ID, ChatState(last_seen_id=5, first_unanswered_id=5, first_unanswered_ts=1))
        missed = Counter()
        _reconcile(self.state, "ivan", set(), missed)
        _reconcile(self.state, "ivan", set(), missed)
        _reconcile(self.state, "ivan", {LEAD_ID}, missed)
        _reconcile(self.state, "ivan", set(), missed)
        _reconcile(self.state, "ivan", set(), missed)
        self.assertIsNotNone(self.state.get("ivan", LEAD_ID))
        with self.assertLogs("sla.monitor", "INFO"):
            _reconcile(self.state, "ivan", set(), missed)
        self.assertIsNone(self.state.get("ivan", LEAD_ID))

    async def test_orphan_statistics_are_dropped(self):
        self.state.save("ivan", LEAD_ID, ChatState(last_seen_id=9, first_unanswered_id=9, first_unanswered_ts=1))
        self.state.conn.executemany("INSERT INTO stats_answers (session, chat_id, message_id, opened_ts, closed_ts) VALUES ('ivan', ?, ?, 1, NULL)",
                                    [(LEAD_ID, 7), (LEAD_ID, 9)])
        self.state.conn.commit()
        with self.assertLogs("sla.monitor", "INFO"):
            _reconcile(self.state, "ivan", {LEAD_ID}, Counter())
        rows = self.state.conn.execute("SELECT message_id FROM stats_answers").fetchall()
        self.assertEqual(rows, [(9,)])


# ---------------------------------------------------------------- process-level

class InstanceLockTests(unittest.TestCase):
    def test_second_copy_is_refused(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "state.sqlite3.lock"
            first = InstanceLock(path).acquire()
            with self.assertRaises(AlreadyRunning):
                InstanceLock(path).acquire()
            first.release()
            InstanceLock(path).acquire().release()


class StatsBotConflictTests(unittest.IsolatedAsyncioTestCase):
    async def test_409_conflict_goes_to_developer(self):
        from sla.stats_bot import StatsBot
        replies = [{"ok": True}, {"ok": False, "error_code": 409, "description": "Conflict"},
                   {"ok": True, "result": []}]
        bot = DeliveryBot()

        async def request(method, data=None):
            if not replies:
                raise asyncio.CancelledError
            return replies.pop(0)
        bot.request = request

        async def no_wait(_):
            return None
        tech = TechAlerts(dev_cfg(), bot)
        with patch("sla.stats_bot.asyncio.sleep", no_wait), self.assertLogs("sla.stats_bot", "WARNING"):
            with self.assertRaises(asyncio.CancelledError):
                await StatsBot(dev_cfg(sessions=[]), None, bot, tech).poll()
        self.assertEqual(len(bot.to(DEV)), 2)
        self.assertIn("409", bot.to(DEV)[0])
        self.assertIn("закончился", bot.to(DEV)[1])
        self.assertEqual(bot.to(MAIN), [])


class RunLifecycleTests(unittest.IsolatedAsyncioTestCase):
    async def test_start_and_stop_messages_only_to_developer(self):
        import sla_monitor
        import sla.monitor as monitor_module
        import sla.notifier as notifier_module
        bot = DeliveryBot()

        class FakeNotifier(DeliveryBot):
            def __init__(self, cfg):
                pass
            async def start(self):
                pass
            async def stop(self):
                pass
            async def deliver(self, recipient, body):
                return await bot.deliver(recipient, body)

        async def quiet_session(*args, **kwargs):
            return None
        with tempfile.TemporaryDirectory() as td:
            cfg = dict(dev_cfg(stats_enabled=False, edit_resolved_alerts=False, history_max_age_hours=72),
                       state_db=str(Path(td) / "state.sqlite3"), sessions=[dict(OP)])
            with patch.object(notifier_module, "Notifier", FakeNotifier), \
                    patch.object(monitor_module, "monitor_session", quiet_session):
                await sla_monitor.run(cfg)
            self.assertFalse(os.path.exists(Path(td) / "missing"))
        dev = bot.to(DEV)
        self.assertEqual(len(dev), 2)
        self.assertIn("запущен", dev[0])
        self.assertIn("остановлен", dev[1])
        self.assertEqual([who for who, _ in bot.sent], [DEV, DEV])


class ReviewFindingsTests(unittest.IsolatedAsyncioTestCase):
    """Regressions for the independent review of v0.5.5."""
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        path = Path(self.tmp.name) / "operator.session"
        path.touch()
        self.sess = dict(OP, session_file=str(path.with_suffix("")))
        self.state = State(Path(self.tmp.name) / "state.sqlite3")
        self.bot = DeliveryBot()
        self.sleeps = []

    async def asyncTearDown(self):
        self.state.close()
        self.tmp.cleanup()

    def cfg(self, **extra):
        return dict(dev_cfg(history_request_delay_seconds=0, poll_interval_seconds=123, **extra),
                    api_id=1, api_hash="a" * 32, include_archived=True, flood_wait_extra_seconds=3)

    async def run_session(self, client_cls, cfg, bot=None):
        bot = bot or self.bot

        async def fake_sleep(seconds):
            self.sleeps.append(seconds)
            if seconds == cfg["poll_interval_seconds"]:
                raise asyncio.CancelledError
        with patch.dict(sys.modules, {"telethon": make_telethon(client_cls)}), \
                patch("sla.monitor.asyncio.sleep", fake_sleep):
            try:
                await monitor_session(cfg, self.sess, self.state, bot, TechAlerts(cfg, bot))
                return "returned"
            except asyncio.CancelledError:
                return "cancelled"

    async def test_request_stuck_after_telethon_gave_up_triggers_reconnect(self):
        class Client:
            instances = 0
            def __init__(self, *args):
                type(self).instances += 1
                self.first = type(self).instances == 1
                self.disconnected = asyncio.get_running_loop().create_future()
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                if self.first:
                    # Queued in Telethon's send queue while its reconnect fails:
                    self.disconnected.set_exception(ConnectionError("Automatic reconnection failed"))
                    await asyncio.get_running_loop().create_future()  # never resolved
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            result = await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0))
        self.assertEqual(result, "cancelled")
        self.assertEqual(Client.instances, 2)
        self.assertIn("мониторинг не работает", self.bot.to(DEV)[0])
        self.assertIn("восстановлена", self.bot.to(DEV)[1])

    async def test_hung_request_without_disconnect_signal_is_detected(self):
        class Client:
            instances = 0
            def __init__(self, *args):
                type(self).instances += 1
                self.first = type(self).instances == 1
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                if self.first:
                    await asyncio.get_running_loop().create_future()
                return
                yield

        with patch("sla.monitor.STALL_LIMIT_SECONDS", 0.2), patch("sla.monitor.WATCH_INTERVAL_SECONDS", 0.05), \
                self.assertLogs("sla.monitor", "WARNING"):
            result = await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0))
        self.assertEqual(result, "cancelled")
        self.assertEqual(Client.instances, 2)
        self.assertEqual(len(self.bot.to(DEV)), 2)

    async def test_telethon_cancelling_its_requests_means_reconnect(self):
        class Client:
            instances = 0
            def __init__(self, *args):
                type(self).instances += 1
                self.first = type(self).instances == 1
                self.connected = True
            def is_connected(self): return self.connected
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): self.connected = False
            async def iter_dialogs(self, archived):
                if self.first:
                    request = asyncio.get_running_loop().create_future()

                    def telethon_disconnects():
                        self.connected = False  # _user_connected = False, then future.cancel()
                        request.cancel()
                    asyncio.get_running_loop().call_later(0.01, telethon_disconnects)
                    await request
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            result = await self.run_session(Client, self.cfg())
        self.assertEqual(result, "cancelled")  # the test's own stop, from the 2nd client
        self.assertEqual(Client.instances, 2)

    async def test_internal_cancel_does_not_stop_the_whole_process(self):
        import sla_monitor
        import sla.notifier as notifier_module
        loop = asyncio.get_running_loop()
        bot = DeliveryBot()

        class FakeNotifier(DeliveryBot):
            def __init__(self, cfg): pass
            async def start(self): pass
            async def stop(self): pass
            async def deliver(self, recipient, body):
                return await bot.deliver(recipient, body)

        class Client:  # no is_connected(): only supervise() can tell it is not a shutdown
            def __init__(self, session_file, *args):
                self.victim = session_file.endswith("victim")
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR if self.victim else 444)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                if self.victim:
                    request = loop.create_future()
                    loop.call_later(0.05, request.cancel)
                    await request
                return
                yield

        with tempfile.TemporaryDirectory() as tmp:
            for name in ("victim", "other"):
                (Path(tmp) / f"{name}.session").touch()
            cfg = dict(dev_cfg(stats_enabled=False, edit_resolved_alerts=False, poll_interval_seconds=10,
                               history_request_delay_seconds=0, include_archived=True),
                       api_id=1, api_hash="a" * 32, state_db=str(Path(tmp) / "state.sqlite3"),
                       sessions=[dict(OP, session_file=str(Path(tmp) / "victim")),
                                 dict(OP, name="petr", operator_label="Пётр", operator_telegram_id=444,
                                      session_file=str(Path(tmp) / "other"))])
            with patch.object(notifier_module, "Notifier", FakeNotifier), \
                    patch.dict(sys.modules, {"telethon": make_telethon(Client)}), \
                    self.assertLogs("sla", "ERROR"):
                task = asyncio.ensure_future(sla_monitor.run(cfg))
                await asyncio.wait({task}, timeout=1.0)
                survived = not task.done()
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)
        self.assertTrue(survived)
        dev = bot.to(DEV)
        self.assertTrue(any("прервана изнутри" in body for body in dev))
        self.assertIn("остановлен", dev[-1])

    async def test_floodwait_raised_by_connect_is_obeyed(self):
        class Client:
            connects = 0
            def __init__(self, *args): pass
            async def connect(self):
                type(self).connects += 1
                if type(self).connects == 1:
                    raise FakeErrors.FloodWaitError(3600)
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            self.assertEqual(await self.run_session(Client, self.cfg()), "cancelled")
        self.assertEqual(self.sleeps[0], 3603)
        self.assertIn("ограничил запросы при подключении", self.bot.to(DEV)[0])
        self.assertNotIn("внутренняя ошибка", " ".join(self.bot.to(DEV)))

    async def test_outage_ending_in_floodwait_resolves_the_link_alert(self):
        class Client:
            connects = 0
            def __init__(self, *args): pass
            async def connect(self):
                type(self).connects += 1
                if type(self).connects == 1:
                    raise ConnectionError("network is unreachable")
                if type(self).connects == 2:
                    raise FakeErrors.FloodWaitError(30)
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        with self.assertLogs("sla.monitor", "WARNING"):
            await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0))
        dev = self.bot.to(DEV)
        self.assertEqual(len(dev), 2)
        self.assertIn("мониторинг не работает", dev[0])
        self.assertIn("восстановлена", dev[1])
        self.assertEqual(self.sleeps[:2], [5, 33])

    async def test_login_that_never_answers_is_bounded(self):
        class Client:
            connects = 0
            def __init__(self, *args): pass
            async def connect(self):
                type(self).connects += 1
                if type(self).connects == 1:  # TCP up, but Telegram never replies
                    await asyncio.get_running_loop().create_future()
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        with patch("sla.monitor.LOGIN_TIMEOUT_SECONDS", 0.1), self.assertLogs("sla.monitor", "WARNING"):
            self.assertEqual(await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0)), "cancelled")
        self.assertEqual(Client.connects, 2)
        self.assertIn("мониторинг не работает", self.bot.to(DEV)[0])

    async def test_whole_host_outage_is_reported_after_recovery(self):
        network = {"up": False}

        class Client:
            connects = 0
            def __init__(self, *args): pass
            async def connect(self):
                type(self).connects += 1
                if type(self).connects == 4:
                    network["up"] = True
                if not network["up"]:
                    raise ConnectionError("network is unreachable")
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                return
                yield

        class Bot(DeliveryBot):
            async def deliver(self, recipient, body):
                if not network["up"]:
                    return Delivery("transient", None, "ClientConnectorError")
                return await super().deliver(recipient, body)

        bot = Bot()
        with self.assertLogs("sla", "WARNING"):
            await self.run_session(Client, self.cfg(tech_alert_delay_seconds=0), bot)
        dev = bot.to(DEV)
        self.assertEqual(len(dev), 1)
        self.assertIn("восстановлена", dev[0])
        self.assertIn("сообщить не удалось", dev[0])
        self.assertIn("мониторинг не работает", dev[0])

    async def test_generic_406_is_not_treated_as_dead_session(self):
        class Client:
            passes = 0
            def __init__(self, *args): pass
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                type(self).passes += 1
                raise FakeErrors.AuthKeyError("FILEREF_UPGRADE_NEEDED")
                yield

        with self.assertLogs("sla.monitor", "ERROR"):
            self.assertEqual(await self.run_session(Client, self.cfg()), "cancelled")
        self.assertEqual(self.bot.to(DEV), [])

        class Duplicated(Client):
            async def iter_dialogs(self, archived):
                raise FakeErrors.AuthKeyDuplicatedError("AUTH_KEY_DUPLICATED")
                yield
        with self.assertLogs("sla.monitor", "ERROR"):
            self.assertEqual(await self.run_session(Duplicated, self.cfg()), "returned")
        self.assertIn("отклонил сессию", self.bot.to(DEV)[0])

    async def test_operator_label_is_escaped_in_developer_message(self):
        cfg = dev_cfg(alert_mode="unanswered")
        bot = DeliveryBot(permanent=[OPERATOR])
        op = dict(OP, operator_label="<Иван & Co>")
        msgs = [FakeMessage(1, 90)]
        await process_dialog(cfg, op, FakeClient(msgs), fake_dialog(msgs), self.state, bot, OPERATOR, BASE,
                             tech=TechAlerts(cfg, bot))
        self.assertIn("&lt;Иван &amp; Co&gt;", bot.to(DEV)[0])


class StatisticsKeepsBreachesTests(StateCase):
    async def test_withdraw_and_forget_keep_overdue_cases(self):
        from sla.statistics import withdraw_case
        now = 10_000
        rows = [("ivan", 1, 10, now - 7200), ("ivan", 1, 11, now - 60),
                ("ivan", 2, 20, now - 7200), ("ivan", 2, 21, now - 60)]
        self.state.conn.executemany("INSERT INTO stats_answers (session, chat_id, message_id, opened_ts, closed_ts) VALUES (?, ?, ?, ?, NULL)", rows)
        self.state.conn.commit()
        withdraw_case(self.state.conn, "answers", "ivan", 1, 10, now, 3600)
        withdraw_case(self.state.conn, "answers", "ivan", 1, 11, now, 3600)
        self.state.forget_chat("ivan", 2, now, 3600)
        left = self.state.conn.execute("SELECT chat_id, message_id, closed_ts FROM stats_answers "
                                       "ORDER BY message_id").fetchall()
        self.assertEqual(left, [(1, 10, now), (2, 20, now)])  # overdue kept as breaches


class SecondReviewTests(StateCase):
    async def test_withdrawn_breach_is_not_an_answer(self):
        from sla.statistics import snapshot, withdraw_case
        from sla.stats_excel import make_excel
        from openpyxl import load_workbook
        import io
        cfg = dict(fake_cfg(), sessions=[dict(OP)], quiet_window=None)
        now = int(BASE.timestamp())
        self.state.conn.execute("INSERT INTO stats_answers (session, chat_id, message_id, opened_ts) "
                                "VALUES ('ivan', 1, 10, ?)", (now - 3 * 3600,))
        self.state.conn.commit()
        withdraw_case(self.state.conn, "answers", "ivan", 1, 10, now, 3600)
        report = snapshot(self.state.conn, cfg, ["ivan"], "today", BASE)
        self.assertEqual(report["answers"]["handled"], 0)
        self.assertIsNone(report["answers"]["avg"])
        self.assertEqual(report["answers"]["breaches"], 1)
        self.assertEqual(report["live"]["answers"], 0)
        _, data = make_excel(cfg, self.state, ["ivan"], "today", now=BASE)
        sheet = load_workbook(io.BytesIO(data))["Ответы"]
        self.assertIn("удалено", sheet.cell(row=2, column=8).value)
        self.assertIsNone(sheet.cell(row=2, column=6).value)

    async def test_long_but_moving_scan_is_not_a_stall(self):
        loop = asyncio.get_running_loop()

        async def pause(seconds):
            future = loop.create_future()
            loop.call_later(seconds, future.set_result, None)
            await future

        class Client:
            instances = 0
            def __init__(self, *args):
                type(self).instances += 1
            async def connect(self): pass
            async def is_user_authorized(self): return True
            async def get_me(self): return SimpleNamespace(id=OPERATOR)
            async def disconnect(self): pass
            async def iter_dialogs(self, archived):
                yield dialog_for(LEAD_ID, [FakeMessage(i, 60 - i) for i in range(1, 11)])
            async def iter_messages(self, entity, **kwargs):
                for i in range(10, 0, -1):
                    await pause(0.05)  # 0.5 s in total, limit is 0.2 s without progress
                    yield FakeMessage(i, 60 - i)

        cfg = dict(fake_cfg(alert_mode="unanswered", history_request_delay_seconds=0, poll_interval_seconds=123,
                            sla_seconds=10 ** 9),
                   api_id=1, api_hash="a" * 32, include_archived=True, flood_wait_extra_seconds=3)
        path = Path(self.tmp.name) / "op.session"
        path.touch()

        async def fake_sleep(seconds):
            if seconds == 123:
                raise asyncio.CancelledError
        with patch.dict(sys.modules, {"telethon": make_telethon(Client)}), \
                patch("sla.monitor.asyncio.sleep", fake_sleep), \
                patch("sla.monitor.STALL_LIMIT_SECONDS", 0.2), patch("sla.monitor.WATCH_INTERVAL_SECONDS", 0.05):
            with self.assertRaises(asyncio.CancelledError):
                await monitor_session(cfg, dict(OP, session_file=str(path.with_suffix(""))), self.state,
                                      DeliveryBot(), TechAlerts(cfg, DeliveryBot()))
        self.assertEqual(Client.instances, 1)  # no false stall / reconnect
        self.assertEqual(self.state.get("ivan", LEAD_ID).first_unanswered_id, 1)

    async def test_outage_is_not_backdated_across_a_known_sleep(self):
        from sla.monitor import _Progress
        import time as time_module
        progress = _Progress()
        progress.beat(idle_seconds=3000)  # e.g. a 50-minute FloodWait just started
        self.assertAlmostEqual(progress.outage_start(), time_module.time(), delta=1)


class DeveloperCommandsTests(StateCase):
    """v0.5.6: every bot command available to the supervisor is available to the developer."""
    async def test_developer_has_all_commands_operator_has_none(self):
        from sla.stats_bot import StatsBot
        from test_statistics import FakeBot, config
        cfg = dict(config(), developer_telegram_id=DEV)
        bot = FakeBot()
        handler = StatsBot(cfg, self.state, bot)
        msg = lambda sender, chat, text, kind="private": {
            "text": text, "from": {"id": sender}, "chat": {"id": chat, "type": kind}}
        callback = lambda sender, data: {"id": "cb", "data": data, "from": {"id": sender},
                                         "message": {"chat": {"id": sender, "type": "private"}, "message_id": 2}}
        # Operator and wrong chats: nothing.
        await handler.on_message(msg(OPERATOR, OPERATOR, "/stats"))
        await handler.on_message(msg(DEV, -100500, "/stats", "supergroup"))
        await handler.on_message(msg(DEV, MAIN, "/stats"))
        self.assertEqual(bot.requests, [])
        # Developer: /stats, every button kind, both Excel variants.
        await handler.on_message(msg(DEV, DEV, "/stats"))
        for data in ("p:week", "o:today:0", "l:a:0", "x:today:0"):
            await handler.on_callback(callback(DEV, data))
        await handler.on_message(msg(DEV, DEV, "/stats_excel"))
        await handler.on_message(msg(DEV, DEV, "/stats_excel ivan 01.09.2026 24.09.2026"))
        self.assertEqual(len(bot.documents), 3)
        self.assertEqual({who for who, *_ in bot.documents}, {DEV})
        self.assertEqual(sum(action == "sendMessage" for action, _ in bot.requests), 1)
        self.assertEqual(sum(action == "editMessageText" for action, _ in bot.requests), 3)
        # An operator pressing a stale button still gets only "Нет доступа".
        bot.requests.clear()
        await handler.on_callback(callback(OPERATOR, "p:today"))
        self.assertEqual([action for action, _ in bot.requests], ["answerCallbackQuery"])

    async def test_without_developer_only_supervisor(self):
        from sla.stats_bot import StatsBot
        from test_statistics import FakeBot, config
        bot = FakeBot()
        handler = StatsBot(config(), self.state, bot)
        await handler.on_message({"text": "/stats", "from": {"id": DEV}, "chat": {"id": DEV, "type": "private"}})
        self.assertEqual(bot.requests, [])


if __name__ == "__main__":
    unittest.main()
